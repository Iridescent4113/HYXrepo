"""Generate the editable Demo plan and two Gantt figures from demo-tasks.csv."""
from pathlib import Path
from datetime import date, timedelta
import csv
import re
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from matplotlib.font_manager import FontProperties, fontManager
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
START, END = date(2026,10,19), date(2026,11,15)
FONTFILE = '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
FP = FontProperties(fname=FONTFILE)
fontManager.addfont(FONTFILE)
plt.rcParams.update({'font.family':FP.get_name(),'axes.unicode_minus':False,'pdf.fonttype':42,'svg.fonttype':'none'})
COLORS = {'产品':'#244F70','后端/数据':'#087F8C','前端':'#269E9B','AI':'#596CB0'}
WEEKS = [
    ('W1','场景、样本与契约',date(2026,10,19),date(2026,10,25),'范围/接口草案冻结，样本与双视图原型走查'),
    ('W2','曲线、画像与基础解释',date(2026,10,26),date(2026,11,1),'样本能导入，曲线与画像可解释'),
    ('W3','角色建议、审核与反馈',date(2026,11,2),date(2026,11,8),'两角色候选→模拟审核→用户反馈'),
    ('W4','联调、验收与演示交接',date(2026,11,9),date(2026,11,15),'稳定演示、降级提示、启动与重置说明'),
]
MILESTONES = [
    ('范围/契约草案冻结',date(2026,10,23),'产品＋后端/数据','一个场景，冻结接口草案'),
    ('样本与原型走查',date(2026,10,25),'产品＋前端＋后端/数据','内置样本与双视图路径可走查'),
    ('数据与画像验收',date(2026,11,1),'产品＋前端＋后端/数据＋AI','曲线、统计、画像与基础解释可演示'),
    ('审核反馈流程验收',date(2026,11,8),'产品＋全员','建议经过模拟审核并记录反馈'),
    ('Demo交付',date(2026,11,15),'产品＋全员','稳定闭环、演示脚本与已知限制'),
]

def tasks():
    with (ROOT/'demo-tasks.csv').open(encoding='utf-8',newline='') as f:rs=list(csv.DictReader(f))
    assert len(rs)==16 and sum(int(r['权重百分比']) for r in rs)==100
    by_id={r['ID']:r for r in rs}
    for r in rs:
        r['start']=date.fromisoformat(r['计划开始']);r['end']=date.fromisoformat(r['计划完成'])
        assert START<=r['start']<=r['end']<=END
    for r in rs:
        for segment in r['依赖'].split(';'):
            if segment=='无':continue
            m=re.fullmatch(r'([D0-9/]+)\((FS|SS)\)',segment);assert m,segment
            for key in m[1].split('/'):
                p=by_id[key]
                assert r['start']>p['end'] if m[2]=='FS' else r['start']>=p['start'],(r['ID'],key)
    return rs

def scaffold(ws,title,note,lastcol):
    ws.sheet_view.showGridLines=False
    for row,text in [(1,title),(2,note)]:
        ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=lastcol)
        ws.cell(row,1,text).font=Font(name='Microsoft YaHei',size=18 if row==1 else 10,bold=row==1,color='132E43')
        ws.cell(row,1).alignment=Alignment(wrap_text=True,vertical='center')
        ws.row_dimensions[row].height=34 if row==1 else 32
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1
    ws.oddFooter.center.text='AI＋代谢Demo | 计划基线 | &P'

def header(ws,values):
    for c,v in enumerate(values,1):
        x=ws.cell(5,c,v);x.fill=PatternFill('solid',fgColor='132E43')
        x.font=Font(name='Microsoft YaHei',color='FFFFFF',bold=True,size=10)
        x.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    ws.row_dimensions[5].height=32

def workbook(rs):
    w=Workbook();tl=w.active;tl.title='任务时间线';team=w.create_sheet('角色分工');weekly=w.create_sheet('四周计划');ms=w.create_sheet('里程碑');helpws=w.create_sheet('使用说明')
    w.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True,calcMode='auto')
    scaffold(tl,'AI＋代谢Demo｜时间线与任务分配','2026-10-19～11-15，共4周；16项任务。负责人待具名，实际实现进度从0%起。',16)
    header(tl,['ID','任务','主责角色','协作角色','计划开始','计划完成','天数','权重','证据成熟度','实际贡献','依赖','验收证据']+[f'W{k+1}' for k in range(4)])
    for k in range(4):
        tl.cell(4,13+k,START+timedelta(weeks=k)).number_format='mm-dd'
    for row,r in enumerate(rs,6):
        vals=[r['ID'],r['任务'],r['主责角色'],r['协作角色'],r['start'],r['end'],f'=F{row}-E{row}+1',int(r['权重百分比'])/100,0,f'=H{row}*I{row}',r['依赖'],r['验收证据']]
        for c,v in enumerate(vals,1):tl.cell(row,c,v)
        for c in [5,6]:tl.cell(row,c).number_format='yyyy-mm-dd';tl.cell(row,c).fill=PatternFill('solid',fgColor='E5F3F5')
        for c in [8,9,10]:tl.cell(row,c).number_format='0%'
        for c in range(13,17):
            letter=get_column_letter(c)
            tl.cell(row,c,f'=IF(AND($E{row}<={letter}$4+6,$F{row}>={letter}$4),"■","")')
            tl.cell(row,c).font=Font(name='Arial',color='FFFFFF',size=14)
            tl.cell(row,c).alignment=Alignment(horizontal='center',vertical='center')
    for role,paint in COLORS.items():
        condition=f'AND($C6="{role}",$E6<=M$4+6,$F6>=M$4)'
        tl.conditional_formatting.add('M6:P21',FormulaRule(formula=[condition],fill=PatternFill('solid',fgColor=paint[1:])))
    tl['G23']='合计';tl['H23']='=SUM(H6:H21)';tl['J23']='=SUM(J6:J21)'
    tl['H23'].number_format=tl['J23'].number_format='0%'
    dv=DataValidation(type='list',formula1='"0,0.25,0.5,0.75,1"',allow_blank=False)
    dv.showErrorMessage=True;dv.error='请选择0、0.25、0.5、0.75、1，且需保留验收依据。';tl.add_data_validation(dv);dv.add('I6:I21')
    for c,width in {'A':8,'B':35,'C':16,'D':27,'E':14,'F':14,'G':8,'H':9,'I':13,'J':13,'K':29,'L':60,'M':10,'N':10,'O':10,'P':10}.items():tl.column_dimensions[c].width=width
    tl.freeze_panes='M6';tl.auto_filter.ref='A5:P21';tl.print_title_rows='1:5'
    scaffold(team,'Demo角色分工｜3–4人，按岗位认领','临床顾问是内容支持角色；本表是建议分工，不是已经落实的人员配置。',5)
    header(team,['角色','建议配置','主要工作','可兼任方式','任务ID'])
    entries=[
        ('项目/产品','1人','场景范围、原型协调、风险边界、验收与演示','兼任设计/测试组织','D01/D13/D15'),
        ('前端','1人','曲线、画像、建议、模拟审核与反馈页面','同一网页覆盖患者/医生视图','D04/D06/D08/D12'),
        ('后端/数据','1人','样本导入、统计存储、审核接口、联调与部署','合并数据与基础设施职责','D02/D03/D05/D07/D11/D14/D16'),
        ('AI','1人或兼任','解释适配器、饮食/活动候选、合并与降级','3人时可由有能力的后端兼任，减少可选LLM联调','D09/D10'),
        ('临床顾问','内容支持','审核演示内容、适用条件和禁忌边界','无人可审核时仅用一般教育与模板','协作D10/D13'),
    ]
    for row,vals in enumerate(entries,6):
        for c,v in enumerate(vals,1):team.cell(row,c,v)
    for c,width in {'A':22,'B':20,'C':70,'D':70,'E':45}.items():team.column_dimensions[c].width=width
    scaffold(weekly,'四周项目计划','默认2026-10-19启动；周末里程碑是计划，不代表已完成。',6)
    header(weekly,['周次','重点','计划开始','计划完成','周末验收','协作范围'])
    for row,(week,name,a,b,accept) in enumerate(WEEKS,6):
        for c,v in enumerate([week,name,a,b,accept,'产品＋前端＋后端/数据＋AI'],1):weekly.cell(row,c,v)
        weekly.cell(row,3).number_format=weekly.cell(row,4).number_format='yyyy-mm-dd'
    for c,width in {'A':10,'B':34,'C':16,'D':16,'E':75,'F':40}.items():weekly.column_dimensions[c].width=width
    scaffold(ms,'Demo轻量里程碑','本期使用Demo验收点；真实试点再进入原IPD与适用机构准入流程。',4)
    header(ms,['里程碑','目标日期','牵头角色','验收证据'])
    for row,vals in enumerate(MILESTONES,6):
        for c,v in enumerate(vals,1):ms.cell(row,c,v)
        ms.cell(row,2).number_format='yyyy-mm-dd'
    for c,width in {'A':28,'B':18,'C':45,'D':90}.items():ms.column_dimensions[c].width=width
    helpws.append(['说明','内容'])
    for pair in [
        ('版本与范围','Demo v0.1；2026-10-13编制。仅设计和计划，尚未开发运行。正式项目原排期不受本文件影响。'),
        ('时间轴','4周，2026-10-19～11-15，起止日计入时长。Excel按周示意；PNG/SVG/PDF按日定位。'),
        ('修改日期','修改任务表E/F列；公式自动计算天数和周条带。超出4周时需扩展列或重新基线化。首次打开请自动重算。'),
        ('进度','I列0/0.25/0.5/0.75/1分别对应未启动/方案评审/模块验收/集成/目标场景验收；纯文档采用专属验收状态。J23为总实际贡献，当前0%。'),
        ('依赖含义','FS：前项完成后才能开始；SS：前项开始后可并行，先用冻结接口/样本联调。完整集成仍需前项达到交付门槛。'),
        ('资源','3–4人是建议。若后端兼任AI，应优先模板模式并检查工作量；本文件未进行资源约束排程或关键路径计算。'),
        ('数据与AI','默认使用模拟样本或授权去标识数据。模板模式无密钥可演示；LLM联调可选并标明模式，不宣称自有模型已训练。'),
        ('医生视图','模拟审核与视图切换不代表真实医生身份验证或正式医疗准入。Demo不用于诊疗。'),
        ('更新静态图','编辑Excel不会更新PNG/PDF；先维护demo-tasks.csv，再运行build_demo_plan.py。运行会覆盖Excel，保留手工改动后再执行。'),
    ]:helpws.append(pair)
    helpws.column_dimensions['A'].width=22;helpws.column_dimensions['B'].width=125
    for ws in w:
        for row in ws.iter_rows():
            for cell in row:
                if cell.row>=6 or ws==helpws:
                    if ws==tl and cell.column>=13:continue
                    cell.font=Font(name='Microsoft YaHei',size=10,color='132E43')
                    cell.alignment=Alignment(wrap_text=True,vertical='center')
            if row[0].row>=6 or ws==helpws:ws.row_dimensions[row[0].row].height=44
        ws.print_area=f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'
    p=ROOT/'AI-metabolic-demo-plan.xlsx';w.save(p)
    check=load_workbook(p,data_only=False)
    assert len(check.sheetnames)==5
    for row,r in enumerate(rs,6):
        assert check['任务时间线'].cell(row,5).value.date()==r['start']
        assert check['任务时间线'].cell(row,6).value.date()==r['end']
    assert abs(sum(check['任务时间线'].cell(row,8).value for row in range(6,22))-1)<1e-10

def frame(fig,ax,title,subtitle,n,left):
    fig.patch.set_facecolor('#F4F7FA');ax.set_facecolor('#FFFFFF')
    fig.subplots_adjust(left=left,right=0.975,top=0.81,bottom=0.15)
    fig.text(.03,.945,title,fontproperties=FP,fontsize=23,color='#132E43')
    fig.text(.03,.895,subtitle,fontproperties=FP,fontsize=11,color='#5A6E7D')
    ax.set_xlim(mdates.date2num(START),mdates.date2num(END+timedelta(days=1)));ax.set_ylim(n-.5,-.5);ax.set_yticks([])
    for k in range(4):
        a=START+timedelta(weeks=k)
        if k%2==0:ax.axvspan(mdates.date2num(a),mdates.date2num(a+timedelta(days=7)),color='#F4F7FA')
        ax.axvline(mdates.date2num(a),color='#D9E3EA',linewidth=.8)
    ax.set_xticks([mdates.date2num(START+timedelta(weeks=k)) for k in range(4)])
    ax.set_xticklabels([f'W{k+1}\n{START+timedelta(weeks=k):%m-%d}' for k in range(4)],fontsize=11)
    ax.xaxis.tick_top();ax.tick_params(axis='x',length=0,pad=9)
    for y in range(n):ax.axhline(y+.5,color='#D9E3EA',linewidth=.6)
    for sp in ax.spines.values():sp.set_visible(False)

def plots(rs):
    fig,ax=plt.subplots(figsize=(21,8));frame(fig,ax,'AI＋代谢Demo｜4周阶段甘特图','一个场景 · 两个智能体角色 · 模拟医生审核与反馈 · 计划基线，当前实现进度0%',5,.35)
    trans=ax.get_yaxis_transform()
    for y,(week,name,a,b,accept) in enumerate(WEEKS):
        ax.barh(y,7,left=mdates.date2num(a),height=.46,color=list(COLORS.values())[y])
        ax.text(-.48,y-.12,f'{week}  {name}',transform=trans,fontproperties=FP,fontsize=13,color='#132E43',va='center',clip_on=False)
        ax.text(-.48,y+.23,f'{a:%m-%d}～{b:%m-%d}',transform=trans,fontproperties=FP,fontsize=10,color='#5A6E7D',va='center',clip_on=False)
        ax.text(mdates.date2num(a)+3.5,y,'1周',fontproperties=FP,fontsize=12,color='white',ha='center',va='center')
    ax.text(-.48,4,'周末验收',transform=trans,fontproperties=FP,fontsize=13,color='#132E43',va='center',clip_on=False)
    for k,label in enumerate(['样本/原型','曲线/画像','审核/反馈','Demo交付']):
        x=mdates.date2num(START+timedelta(days=6+7*k))+.5
        ax.scatter(x,4,marker='D',s=70,color='#A65B13')
        ax.annotate(label,(x,4),xytext=(0,23),textcoords='offset points',fontproperties=FP,fontsize=10,ha='right' if k==3 else 'center',color='#A65B13')
    fig.text(.03,.072,'2026-10-19启动 → 2026-11-15交付；10-23先冻结范围与接口草案。原正式项目IPD排期保持独立。',fontproperties=FP,fontsize=11,color='#132E43')
    fig.text(.03,.04,'默认离线模板可演示；真实CGM与自有模型只预留接口。演示不用于诊疗，不表示临床或模型效果已验证。',fontproperties=FP,fontsize=10,color='#5A6E7D')
    results=[('AI-metabolic-demo-gantt-overview',fig)]
    fig,ax=plt.subplots(figsize=(23,13));frame(fig,ax,'AI＋代谢Demo｜16项任务甘特图','颜色表示主责角色；条带重叠表示并行开发/联调。所有日期为拟议计划，负责人待具名。',16,.39)
    trans=ax.get_yaxis_transform()
    for y,r in enumerate(rs):
        ax.barh(y,(r['end']-r['start']).days+1,left=mdates.date2num(r['start']),height=.48,color=COLORS[r['主责角色']])
        ax.text(-.60,y-.10,f"{r['ID']}  {r['任务']}",transform=trans,fontproperties=FP,fontsize=11,color='#132E43',va='center',clip_on=False)
        ax.text(-.60,y+.25,f"{r['主责角色']}  |  {r['start']:%m-%d}～{r['end']:%m-%d}",transform=trans,fontproperties=FP,fontsize=9,color='#5A6E7D',va='center',clip_on=False)
    for _,d,_,_ in MILESTONES:ax.axvline(mdates.date2num(d)+.5,color='#A65B13',linestyle='--',linewidth=.8,alpha=.8)
    fig.legend([plt.Rectangle((0,0),1,1,color=c) for c in COLORS.values()],COLORS.keys(),loc='lower center',bbox_to_anchor=(.5,.085),ncol=4,prop=FontProperties(fname=FONTFILE,size=11),frameon=False)
    fig.text(.03,.06,'虚线：10-23契约草案冻结、10-25样本原型、11-01数据画像、11-08审核反馈、11-15Demo交付。',fontproperties=FP,fontsize=10,color='#132E43')
    fig.text(.03,.033,'3人时后端兼任AI需检查容量；此图未计算资源约束关键路径。完整依赖、协作、验收和实际进度见Excel。',fontproperties=FP,fontsize=10,color='#5A6E7D')
    results.append(('AI-metabolic-demo-gantt-tasks',fig))
    with PdfPages(ROOT/'AI-metabolic-demo-gantt.pdf') as pdf:
        for name,figure in results:
            figure.savefig(ROOT/f'{name}.png',dpi=160);figure.savefig(ROOT/f'{name}.svg');pdf.savefig(figure);plt.close(figure)

def main():
    rs=tasks();workbook(rs);plots(rs)
    print('Validated: 16 tasks; 100% total weight; 4 weeks; all FS/SS dependencies satisfy scheduled dates; 5 Excel sheets.')

if __name__=='__main__':main()
