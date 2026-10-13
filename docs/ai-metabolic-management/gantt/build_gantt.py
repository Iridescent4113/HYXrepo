"""Build the project Gantt from the existing WBS, without changing its dates."""
from pathlib import Path
from datetime import date, datetime, timedelta
import csv
import zipfile
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from matplotlib.font_manager import FontProperties, fontManager
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT.parent / 'ipd-deliverables.csv'
START, END = date(2026, 10, 19), date(2027, 1, 24)
PHASES = [
    ('概念与资源', '2026-10-19', '2026-11-01', 10),
    ('计划与方案冻结', '2026-11-02', '2026-11-15', 25),
    ('开发与集成', '2026-11-16', '2026-12-13', 60),
    ('验证与试点准入', '2026-12-14', '2026-12-27', 80),
    ('首批试点启动', '2026-12-28', '2027-01-10', 90),
    ('首轮观察与复盘', '2027-01-11', '2027-01-24', 100),
]
MILESTONES = [
    ('DCP1 · 立项', date(2026, 11, 1)),
    ('DCP2 · 冻结', date(2026, 11, 15)),
    ('DCP3 · 准入', date(2026, 12, 27)),
    ('试点启动截止', date(2027, 1, 10)),
    ('DCP4 · 复盘', date(2027, 1, 24)),
]
COLORS = {'需求与立项': '#244F70', '架构与治理': '#527A97',
          '数据采集': '#087F8C', '决策模型': '#4D6AB1',
          '用户与干预': '#219C91', '验证': '#B87B39', '试点运营': '#697681'}
FONTFILE = '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
F = FontProperties(fname=FONTFILE)
fontManager.addfont(FONTFILE)
plt.rcParams.update({'font.family': F.get_name(), 'axes.unicode_minus': False,
                     'pdf.fonttype': 42, 'svg.fonttype': 'none'})

def read_rows():
    with SOURCE.open(encoding='utf-8-sig', newline='') as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 20
    assert sum(float(r['权重百分比']) for r in rows) == 100
    for r in rows:
        r['start'] = date.fromisoformat(r['计划开始'])
        r['end'] = date.fromisoformat(r['计划完成'])
        assert START <= r['start'] <= r['end'] <= END
    return rows

def frame(fig, ax, title, subtitle, count, left=0.34):
    fig.patch.set_facecolor('#F4F7FA'); ax.set_facecolor('#FFFFFF')
    fig.subplots_adjust(left=left, right=0.975, top=0.80, bottom=0.15)
    fig.text(0.03, 0.94, title, fontproperties=F, fontsize=23, color='#132E43', weight='bold')
    fig.text(0.03, 0.895, subtitle, fontproperties=F, fontsize=11, color='#5A6E7D')
    ax.set_xlim(mdates.date2num(START), mdates.date2num(END + timedelta(days=1)))
    ax.set_ylim(count - 0.5, -0.5)
    ax.set_yticks([])
    weeks = [START + timedelta(weeks=k) for k in range(14)]
    ax.set_xticks([mdates.date2num(d) for d in weeks])
    ax.set_xticklabels([f'W{k+1}\n{d:%m-%d}' for k, d in enumerate(weeks)], fontsize=9)
    ax.xaxis.tick_top(); ax.tick_params(axis='x', length=0, pad=9)
    for k, d in enumerate(weeks):
        if k % 2 == 0:
            ax.axvspan(mdates.date2num(d), mdates.date2num(d + timedelta(days=7)), color='#F4F7FA', zorder=0)
        ax.axvline(mdates.date2num(d), color='#D9E3EA', linewidth=0.7, zorder=1)
    for y in range(count): ax.axhline(y+0.5, color='#E3E9EE', linewidth=0.7, zorder=1)
    for s in ax.spines.values(): s.set_visible(False)
    ax.set_axisbelow(True)
    fig.text(left, 0.845, '2026-10-19 → 2027-01-24 / 日期按天定位，结束日计入任务时长',
             fontproperties=F, fontsize=10, color='#132E43')

def overview():
    fig, ax = plt.subplots(figsize=(22, 9))
    frame(fig, ax, 'AI＋代谢管理｜IPD阶段甘特图',
          '条件性计划基线 · 中国大陆 / 医疗机构参与 · 当前产品实际完成度0% · 12周内启动受控试点', 7)
    trans = ax.get_yaxis_transform()
    for y, (name, a, b, target) in enumerate(PHASES):
        start, end = date.fromisoformat(a), date.fromisoformat(b)
        x, days = mdates.date2num(start), (end-start).days+1
        ax.barh(y, days, left=x, height=0.46, color=['#244F70','#527A97','#087F8C','#B87B39','#219C91','#697681'][y], zorder=3)
        ax.text(-0.50, y-0.07, name, transform=trans, fontproperties=F, fontsize=13,
                color='#132E43', va='center', clip_on=False)
        ax.text(-0.50, y+0.22, f'{start:%Y-%m-%d}～{end:%Y-%m-%d}  |  计划累计 {target}%',
                transform=trans, fontproperties=F, fontsize=9.5, color='#5A6E7D', va='center', clip_on=False)
        ax.text(x+days/2, y, f'{days//7}周', fontproperties=F, fontsize=11, color='white', ha='center', va='center')
    y = 6
    ax.text(-0.50, y, '决策 / 截止里程碑', transform=trans, fontproperties=F, fontsize=13, color='#132E43', va='center', clip_on=False)
    for label, d in MILESTONES:
        x=mdates.date2num(d)+0.5
        ax.scatter(x,y,marker='D',s=60,color='#A65B13',zorder=4)
        # The final event label stays inside the chart, others remain centered.
        align='right' if d==END else 'center'
        ax.annotate(f'{label}\n{d:%m-%d}',(x,y),xytext=(0,24),textcoords='offset points',
                    fontproperties=F,fontsize=9,color='#A65B13',ha=align,va='bottom')
    fig.text(0.03, 0.073, 'DCP按交付证据决定继续、附条件继续、暂停或终止；日期到达不等于自动过关。', fontproperties=F,fontsize=11,color='#132E43')
    fig.text(0.03, 0.04, '计划100%仅指首期试点范围完成；不代表公开上市或疗效证明。后入组用户的随访可能超过2027-01-24。',fontproperties=F,fontsize=10,color='#5A6E7D')
    return fig

def detail(rows):
    fig, ax = plt.subplots(figsize=(24, 16))
    frame(fig, ax, 'AI＋代谢管理｜20项交付任务甘特图',
          '与已归档WBS日期逐项一致 · 岗位负责人待具名 · 训练与患者入组均有前置条件',len(rows),left=0.41)
    trans=ax.get_yaxis_transform()
    for y,r in enumerate(rows):
        x=mdates.date2num(r['start']);days=(r['end']-r['start']).days+1
        ax.barh(y,days,left=x,height=0.48,color=COLORS[r['工作包']],zorder=3)
        ax.text(-0.655,y-0.10,f"{r['编号']}  {r['交付物']}",transform=trans,fontproperties=F,
                fontsize=10.5,color='#132E43',va='center',clip_on=False)
        ax.text(-0.655,y+0.23,f"{r['负责人角色（待具名）']}  |  {r['start']:%m-%d}～{r['end']:%m-%d}",
                transform=trans,fontproperties=F,fontsize=8.5,color='#5A6E7D',va='center',clip_on=False)
    for label,d in MILESTONES:
        ax.axvline(mdates.date2num(d)+0.5,color='#A65B13',linestyle='--',linewidth=0.9,zorder=2,alpha=0.8)
    handles=[plt.Rectangle((0,0),1,1,color=c) for c in COLORS.values()]
    fig.legend(handles,COLORS.keys(),loc='lower center',bbox_to_anchor=(0.5,0.084),ncol=7,
               prop=FontProperties(fname=FONTFILE,size=10),frameon=False)
    fig.text(0.03,0.06,'重点依赖：CGM接口与授权、训练数据许可/标签、团队与复核容量、机构及适用准入。模型未达门槛仅离线/影子运行。',fontproperties=F,fontsize=10,color='#132E43')
    fig.text(0.03,0.035,'虚线对应11-01 DCP1、11-15 DCP2、12-27 DCP3、2027-01-10启动截止、01-24 DCP4；本图未计算资源约束关键路径。',fontproperties=F,fontsize=9,color='#5A6E7D')
    return fig

NAVY='132E43'; TEAL='087F8C'; MUTED='5A6E7D'; GRID='D9E3EA'
def setup_sheet(ws, title, note, cols):
    ws.sheet_view.showGridLines=False
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=cols)
    ws.cell(1,1,title).font=Font(name='Microsoft YaHei',size=18,bold=True,color=NAVY)
    ws.row_dimensions[1].height=32
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=cols)
    ws.cell(2,1,note).font=Font(name='Microsoft YaHei',size=10,color=MUTED)
    ws.cell(2,1).alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[2].height=32
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1
    ws.print_options.horizontalCentered=True
    ws.oddFooter.center.text='AI＋代谢管理 | 条件性计划 | 第 &P 页'

def style_header(ws,row,cols):
    for cell in ws[row][:cols]:
        cell.fill=PatternFill('solid',fgColor=NAVY);cell.font=Font(name='Microsoft YaHei',color='FFFFFF',bold=True,size=10)
        cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    ws.row_dimensions[row].height=31

def gantt_cells(ws,first,last,week_col,category_col=None,color=TEAL):
    # Start/end fields E/F are intentionally shared by both sheet layouts.
    for row in range(first,last+1):
        for col in range(week_col,week_col+14):
            letter=get_column_letter(col)
            cell=ws.cell(row,col,f'=IF(AND($E{row}<={letter}$5+6,$F{row}>={letter}$5),"■","")')
            cell.alignment=Alignment(horizontal='center',vertical='center')
            cell.font=Font(name='Arial',color='FFFFFF',size=12)
            cell.border=Border(bottom=Side(style='hair',color=GRID),left=Side(style='hair',color=GRID))
    region=f'{get_column_letter(week_col)}{first}:{get_column_letter(week_col+13)}{last}'
    col=get_column_letter(week_col)
    if category_col:
        for category,paint in COLORS.items():
            condition=f'AND(${category_col}{first}="{category}",$E{first}<={col}$5+6,$F{first}>={col}$5)'
            ws.conditional_formatting.add(region,FormulaRule(formula=[condition],fill=PatternFill('solid',fgColor=paint[1:])))
    else:
        condition=f'AND($E{first}<={col}$5+6,$F{first}>={col}$5)'
        ws.conditional_formatting.add(region,FormulaRule(formula=[condition],fill=PatternFill('solid',fgColor=color)))

def workbook(rows):
    wb=Workbook();overview_ws=wb.active;overview_ws.title='阶段甘特图'
    task_ws=wb.create_sheet('任务甘特图');milestone_ws=wb.create_sheet('里程碑');info=wb.create_sheet('使用说明')
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True,calcMode='auto')
    setup_sheet(overview_ws,'AI＋代谢管理｜阶段甘特图','基线：2026-10-19～2027-01-24；未来完成度为计划值。修改E/F日期后，周条带自动更新。',22)
    heads=['阶段','周次','决策/检查点','计划累计完成度','计划开始','计划完成','天数','阶段交付']
    for c,v in enumerate(heads,1):overview_ws.cell(5,c,v)
    for k in range(14):
        c=9+k;overview_ws.cell(4,c,f'W{k+1}');overview_ws.cell(5,c,START+timedelta(weeks=k)).number_format='mm-dd'
    labels=['DCP1','DCP2','TR2/3','DCP3','试点启动截止','DCP4']
    spans=['W1–2','W3–4','W5–8','W9–10','W11–12','W13–14']
    delivers=['资源/范围/接口路径','PRD/PoC/验收冻结','数据/双端/画像/影子模型','质量/临床/适用准入','首批入组与服务验证','首轮观察/问题关闭/扩展决策']
    for row,(name,a,b,target) in enumerate(PHASES,6):
        vals=[name,spans[row-6],labels[row-6],target/100,date.fromisoformat(a),date.fromisoformat(b),f'=F{row}-E{row}+1',delivers[row-6]]
        for c,v in enumerate(vals,1):overview_ws.cell(row,c,v)
        overview_ws.cell(row,4).number_format='0%'
    for c,w in {'A':21,'B':10,'C':19,'D':17,'E':14,'F':14,'G':8,'H':38}.items():overview_ws.column_dimensions[c].width=w
    for c in range(9,23):overview_ws.column_dimensions[get_column_letter(c)].width=8
    style_header(overview_ws,5,22);gantt_cells(overview_ws,6,11,9)
    overview_ws.freeze_panes='I6';overview_ws.print_title_rows='1:5'
    overview_ws.merge_cells('A14:V14');overview_ws['A14']='每格代表一周；与任务日期有交集则填色。精确起止日期见E/F列；PNG/PDF按天定位。'
    overview_ws.merge_cells('A15:V15');overview_ws['A15']='所有阶段均受资源、接口、数据及适用准入约束；计划100%仅指首期试点范围。'
    setup_sheet(task_ws,'AI＋代谢管理｜20项任务甘特图','日期直接来自已确认WBS；责任人待具名。蓝色日期列可编辑；成熟度须有验收证据，当前均为0。',26)
    headers=['编号','工作包','交付物','负责角色（待具名）','计划开始','计划完成','天数','权重','证据成熟度','实际完成贡献','状态','验收证据']
    for c,v in enumerate(headers,1):task_ws.cell(5,c,v)
    for k in range(14):
        c=13+k;task_ws.cell(4,c,f'W{k+1}');task_ws.cell(5,c,START+timedelta(weeks=k)).number_format='mm-dd'
    for row,r in enumerate(rows,6):
        vals=[r['编号'],r['工作包'],r['交付物'],r['负责人角色（待具名）'],r['start'],r['end'],f'=F{row}-E{row}+1',float(r['权重百分比'])/100,float(r['当前成熟度']),f'=H{row}*I{row}',r['当前状态'],r['验收证据']]
        for c,v in enumerate(vals,1):task_ws.cell(row,c,v)
        for c in [8,9,10]:task_ws.cell(row,c).number_format='0%'
    cols={'A':8,'B':15,'C':43,'D':26,'E':14,'F':14,'G':8,'H':9,'I':13,'J':16,'K':30,'L':65}
    for c,w in cols.items():task_ws.column_dimensions[c].width=w
    for c in range(13,27):task_ws.column_dimensions[get_column_letter(c)].width=8
    style_header(task_ws,5,26);gantt_cells(task_ws,6,25,13,'B')
    task_ws.freeze_panes='M6';task_ws.auto_filter.ref='A5:Z25';task_ws.print_title_rows='1:5'
    task_ws.cell(27,7,'合计');task_ws.cell(27,8,'=SUM(H6:H25)');task_ws.cell(27,10,'=SUM(J6:J25)')
    task_ws.cell(27,8).number_format=task_ws.cell(27,10).number_format='0%'
    maturity=DataValidation(type='list',formula1='"0,0.25,0.5,0.75,1"',allow_blank=False)
    maturity.errorTitle='成熟度无效';maturity.error='请选择0、0.25、0.5、0.75或1，且保留验收证据。';maturity.showErrorMessage=True
    task_ws.add_data_validation(maturity);maturity.add('I6:I25')
    setup_sheet(milestone_ws,'IPD里程碑与试点截止','里程碑是拟议评审/截止日期，不是已完成记录。实际批准须保留机构与项目决议。',5)
    for c,v in enumerate(['里程碑','目标日期','决定/条件','牵头角色','当前状态'],1):milestone_ws.cell(5,c,v)
    conditions=['资源、负责人、范围与合法接入路径落实','方案/PoC/数据用途明确、验收门槛冻结','关键缺陷关闭、适用准入与临床运营容量满足','获准首批入组并运行闭环','基于试点报告决定扩大、延长、调整或停止']
    owners=['IPMT / PDT','IPMT / 技术 / 临床','IPMT / 临床 / QA合规','运营 / PDT / 临床','IPMT / PDT / 临床']
    for row,((label,day),condition,owner) in enumerate(zip(MILESTONES,conditions,owners),6):
        for c,v in enumerate([label,day,condition,owner,'待执行'],1):milestone_ws.cell(row,c,v)
        milestone_ws.cell(row,2).number_format='yyyy-mm-dd'
    style_header(milestone_ws,5,5)
    for col,width in {'A':24,'B':17,'C':70,'D':32,'E':14}.items():milestone_ws.column_dimensions[col].width=width
    info.append(['项目甘特图使用说明','内容'])
    notes=[('来源','project-design.md v0.2 与 ipd-deliverables.csv；20项任务及所有起止日期保持不变。'),('基线','2026-10-19启动；2027-01-10前启动试点；2027-01-24首轮复盘。北京时间。'),('计划与实际','阶段累计完成度为拟议计划；用户确认从零起步，产品实际完成度为0%。'),('如何更新','编辑阶段/任务表的计划开始、计划完成；周条带与时长公式随Excel重算更新。重新运行build_gantt.py会从WBS恢复基线，覆盖此工作簿及静态图。'),('图的精度','Excel每格代表一周，任一日与任务相交即填色；PNG/SVG/PDF条长按实际日历天绘制，起止日均计入。'),('成熟度','0未启动；0.25方案评审；0.5模块验收；0.75集成；1目标场景验收。纯文档采用专属验收状态。'),('贡献计算','每项贡献=权重×证据成熟度；任务表J27汇总实际完成度。不得把计划百分比填成实际成熟度。'),('Excel重算','公式不含预计算缓存，首次打开Excel后自动重算；如未显示色条，设置计算选项为自动。'),('时间轴范围','预置14周。日期超出本基线时，应扩展周列或重新基线化；不会自动生成无限日期列。'),('条件与依赖','CGM接口/授权、训练数据许可/标签、具名团队及临床复核容量、机构与适用审批是前置条件。模型未达门槛仅离线/影子运行。'),('试点范围','14周是首轮计划，后入组用户随访可超出结束日期。100%不代表公众上市或疗效证明。'),('关键路径','当前尚无经确认的资源容量与依赖网络，未进行关键路径法计算。不要把最长条带直接当作已计算关键路径。')]
    for item,value in notes:info.append([item,value])
    info.column_dimensions['A'].width=22;info.column_dimensions['B'].width=120
    for row in info.iter_rows():
        for cell in row:cell.font=Font(name='Microsoft YaHei',size=11);cell.alignment=Alignment(wrap_text=True,vertical='top')
        info.row_dimensions[row[0].row].height=42
    for ws,first,last in [(overview_ws,6,11),(task_ws,6,25),(milestone_ws,6,10)]:
        for row in range(first,last+1):
            ws.row_dimensions[row].height=48 if ws==task_ws else 44
            for cell in ws[row]:
                if cell.column < (13 if ws==task_ws else 9 if ws==overview_ws else 6):
                    cell.font=Font(name='Microsoft YaHei',size=10,color=NAVY)
                    cell.alignment=Alignment(wrap_text=True,vertical='center')
            if ws!=milestone_ws:
                for c in [5,6]:
                    ws.cell(row,c).number_format='yyyy-mm-dd'
                    ws.cell(row,c).fill=PatternFill('solid',fgColor='E5F3F5')
        ws.print_area=f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'
    path=ROOT/'AI-metabolic-IPD-gantt.xlsx';wb.save(path)
    check=load_workbook(path,data_only=False)
    assert len(check.sheetnames)==4
    for n,r in enumerate(rows,6):
        assert check['任务甘特图'].cell(n,5).value.date()==r['start']
        assert check['任务甘特图'].cell(n,6).value.date()==r['end']
        assert check['任务甘特图'].cell(n,10).value==f'=H{n}*I{n}'
    assert abs(sum(check['任务甘特图'].cell(n,8).value for n in range(6,26))-1)<1e-9

def main():
    rows=read_rows();workbook(rows)
    figures=[('AI-metabolic-IPD-gantt-overview',overview()),('AI-metabolic-IPD-gantt-tasks',detail(rows))]
    with PdfPages(ROOT/'AI-metabolic-IPD-gantt.pdf') as pdf:
        for name,fig in figures:
            fig.savefig(ROOT/f'{name}.png',dpi=160)
            fig.savefig(ROOT/f'{name}.svg')
            pdf.savefig(fig)
            plt.close(fig)
    (ROOT/'verification.txt').write_text('20 WBS tasks; weights sum to 100%; all task dates equal source CSV.\n14 calendar weeks, inclusive start/end dates.\nExcel: 4 sheets; date formulas, weekly conditional formats, maturity validation, weighted progress formulas.\nMilestones and phase percentages are conditional targets; actual product baseline=0%.\n',encoding='utf-8')
    print('Built and validated 20-task / 14-week Gantt; source dates unchanged.')

if __name__=='__main__':main()
